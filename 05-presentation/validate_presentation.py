"""Validate the deliverable without exporting any participant values."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
import re
import shutil
import subprocess
import tempfile


def render_check(preview):
    """Render the technical copy and confirm no editable text is lost or clipped."""
    deck = Presentation(preview)
    blocks = []
    for slide in deck.slides:
        for shape in slide.shapes:
            frames = ([shape.text_frame] if shape.has_text_frame
                      else [c.text_frame for row in shape.table.rows for c in row.cells] if shape.has_table
                      else [])
            for tf in frames:
                for paragraph in tf.paragraphs:
                    text = ''.join(r.text for r in paragraph.runs).strip()
                    if text:
                        blocks.append(text)
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['libreoffice', '--headless', '-env:UserInstallation=file://' + tmp + '/profile',
                        '--convert-to', 'pdf', '--outdir', tmp, str(preview)],
                       check=True, capture_output=True, timeout=600)
        pdf = Path(tmp) / (preview.stem + '.pdf')
        # -raw keeps each text object together; -layout interleaves side-by-side columns.
        pages = subprocess.run(['pdftotext', '-raw', str(pdf), '-'],
                               check=True, capture_output=True, text=True).stdout.split('\f')
        if pages and not pages[-1].strip():
            pages.pop()  # pdftotext ends the last page with a trailing form feed
        boxes = subprocess.run(['pdftotext', '-bbox', str(pdf), '-'],
                               check=True, capture_output=True, text=True).stdout
        shutil.copy(pdf, preview.with_suffix('.pdf'))
    flat = re.sub(r'\s+', '', ' '.join(pages))
    issues = [dict(kind='text_missing_from_render', text=b)
              for b in blocks if re.sub(r'\s+', '', b) not in flat]
    # Any rendered word reaching outside the page box would be cut off on screen.
    tree = etree.fromstring(boxes.encode())
    ns = {'x': tree.tag.split('}')[0].strip('{')}
    words = 0
    for n, page in enumerate(tree.iter('{%s}page' % ns['x']), 1):
        pw, ph = float(page.get('width')), float(page.get('height'))
        for w in page.iter('{%s}word' % ns['x']):
            words += 1
            x0, y0, x1, y1 = (float(w.get(k)) for k in ('xMin', 'yMin', 'xMax', 'yMax'))
            if x0 < -.5 or y0 < -.5 or x1 > pw + .5 or y1 > ph + .5:
                issues.append(dict(kind='word_outside_page', page=n, text=w.text,
                                   box=[round(v, 1) for v in (x0, y0, x1, y1)]))
    return {'rendered_pages': len(pages), 'text_blocks_checked': len(blocks),
            'rendered_words_checked': words, 'issues': issues,
            'pdf': str(preview.with_suffix('.pdf'))}

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
SOURCE = REPO / '00-task/assets/ЛЦТ2026 Шаблон презентации.pptx'
OUTPUT = ROOT / 'ЛЦТ2026-задача7-ПРОСВЕТ.pptx'
TEAM = ROOT / 'team-data.md'
CHECKS = ROOT / 'checks'
prs = Presentation(OUTPUT)
template = Presentation(SOURCE)
assert len(prs.slides) == 15
empty = []
fonts = set()
pictures = 0
for index, slide in enumerate(prs.slides, 1):
    assert any(s.has_text_frame and s.text.strip() for s in slide.shapes)
    for shape in slide.shapes:
        if shape.is_placeholder and shape.has_text_frame and not shape.text.strip():
            empty.append((index, shape.shape_id))
        frames = [shape.text_frame] if shape.has_text_frame else [c.text_frame for row in shape.table.rows for c in row.cells] if shape.has_table else []
        for tf in frames:
            for paragraph in tf.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        fonts.add(run.font.name or '(inherited)')
        if shape.shape_type == 13:
            pictures += 1
            with Image.open(BytesIO(shape.image.blob)) as image:
                image.verify()
assert not empty
assert fonts == {'Montserrat'}

# All retained fields of the mandatory block keep their exact template geometry.
allowed_removed = {1: {4}, 2: {2}, 3: {2, 3, 4, 5, 6}, 4: set(), 5: set()}
geometry_checked = 0
for index in range(1, 6):
    src = template.slides[index + 5]
    dst = prs.slides[index - 1]
    current = {s.shape_id: s for s in dst.shapes}
    for shape in src.shapes:
        if shape.shape_id in allowed_removed[index]:
            continue
        target = current[shape.shape_id]
        assert tuple(getattr(shape, k) for k in ('left', 'top', 'width', 'height')) == tuple(getattr(target, k) for k in ('left', 'top', 'width', 'height'))
        geometry_checked += 1

with ZipFile(SOURCE) as original, ZipFile(OUTPUT) as result:
    assert result.testzip() is None
    themes = [f for f in result.namelist() if f.startswith('ppt/theme/') and f.endswith('.xml')]
    assert all(original.read(f) == result.read(f) for f in themes)
    xml = {f: result.read(f) for f in result.namelist() if f.endswith('.xml')}
    for blob in xml.values():
        etree.fromstring(blob)
    for f in result.namelist():
        if f.endswith('.rels'):
            for rel in etree.fromstring(result.read(f)):
                assert not (rel.get('Type', '').endswith('/image') and rel.get('TargetMode') == 'External')

    # Read contact values only in memory and never echo them or put them in metadata.
    contacts = []
    for line in TEAM.read_text().splitlines():
        if line.startswith('|'):
            cells = [x.strip() for x in line.strip('|').split('|')]
            if len(cells) == 6 and cells[0].isdigit():
                contacts.extend(cells[3:5])
    assert len(contacts) == 10
    for value in contacts:
        parts = [name for name, blob in xml.items() if value.encode() in blob]
        assert parts == ['ppt/slides/slide3.xml'], 'Contact is missing or occurs outside the participant slide'

# Compare embedded illustration pixels with the exact selected source regions.
media = json.loads((CHECKS / 'illustrations.json').read_text())
for item in media:
    slide = prs.slides[item['slide'] - 1]
    shape = next(s for s in slide.shapes if s.name == item['label'])
    if item['source'].endswith('.pptx'):
        src_pic = next(s for s in template.slides[5].shapes if s.shape_id == item['shape'])
        assert src_pic.image.blob == shape.image.blob
        continue
    with Image.open(REPO / item['source']) as source:
        expected = source.crop(item['crop']) if item.get('crop') else source.copy()
        expected = expected.convert('RGB')
    with Image.open(BytesIO(shape.image.blob)) as actual:
        actual = actual.convert('RGB')
        assert actual.size == expected.size
        if not item.get('palette'):
            assert actual.tobytes() == expected.tobytes()
            continue
        # Плоская графика пережата фиксированной палитрой: сетка пикселей та же,
        # расходятся только сглаженные края и ровно на записанную при сборке величину.
        deviation = max(high for _, high in ImageChops.difference(actual, expected).getextrema())
        assert deviation == item['max_deviation'] <= 40
        assert len(shape.image.blob) == item['bytes'] < item['source_bytes'] // 10
        assert abs(shape.width / shape.height - actual.width / actual.height) < .01

fit = json.loads((CHECKS / 'fit_check.json').read_text())
assert all(item['fits'] for item in fit)
render = render_check(CHECKS / 'content_preview.pptx')
(CHECKS / 'render_check.json').write_text(json.dumps(render, ensure_ascii=False, indent=2))
assert render['rendered_pages'] == 10 and not render['issues']

summary = {
    'slide_count': len(prs.slides),
    'text_present_on_each_slide': True,
    'original_template_pages': [7, 8, 9, 10, 11],
    'mandatory_original_shape_geometries_checked': geometry_checked,
    'embedded_pictures': pictures,
    'source_image_pixel_comparisons_passed': len(media),
    'palette_packed_illustrations': [{k: item[k] for k in ('label', 'palette', 'max_deviation', 'bytes', 'source_bytes')}
                                     for item in media if item.get('palette')],
    'empty_active_placeholders': [],
    'broken_pictures': [],
    'external_image_links': [],
    'explicit_fonts': sorted(fonts),
    'original_theme_bytes_unchanged': True,
    'all_xml_parsed': True,
    'zip_error': None,
    'contact_fields_only_in_participant_slide': True,
    'text_layout_estimates_checked': len(fit),
    'text_layout_estimate_failures': 0,
    'rendered_technical_pages': render['rendered_pages'],
    'rendered_native_text_blocks_matched': render['text_blocks_checked'],
    'bytes': OUTPUT.stat().st_size,
    'sha256': hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
}
(CHECKS / 'validation.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2))
print(json.dumps(summary, ensure_ascii=False, indent=2))
