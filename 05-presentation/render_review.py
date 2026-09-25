"""Create a local visual-review deck with participant contacts removed.

Input/output are confined to 05-presentation. Values never go to stdout, notes,
metadata or logs. The deliverable itself is not edited. Rendering can be done on
an allowed compute node only after this script has proved redaction succeeded.
"""
from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
source = ROOT / 'ЛЦТ2026-задача7-ПРОСВЕТ.pptx'
output = ROOT / 'checks/review_redacted.pptx'
contacts = []
for line in (ROOT / 'team-data.md').read_text().splitlines():
    if line.startswith('|'):
        cells = [part.strip() for part in line.strip('|').split('|')]
        if len(cells) == 6 and cells[0].isdigit():
            contacts.extend(cells[3:5])
assert len(contacts) == 10
prs = Presentation(source)
assert len(prs.slides) == 16
counts = {value: 0 for value in contacts}
for slide in prs.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                for value in contacts:
                    if value in run.text:
                        counts[value] += 1
                        run.text = run.text.replace(value, 'контакт скрыт')
assert set(counts.values()) == {1}, 'Each contact must occur once in the original slide fields'
output.parent.mkdir(exist_ok=True)
prs.save(output)
with ZipFile(output) as package:
    xml = [package.read(name) for name in package.namelist() if name.endswith('.xml')]
assert not any(value.encode() in part for value in contacts for part in xml), 'Redaction incomplete'
print('PASS: 16 review pages, 10 contact fields redacted; original unchanged')
